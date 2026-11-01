"""Outils contrôlés de Llink Work : recherche web, Python isolé et exports documentaires."""

from __future__ import annotations

import asyncio
import csv
import io
import json
import logging
import os
import re
from typing import Any
from urllib.parse import quote
from xml.sax.saxutils import escape

import httpx
from fastapi import APIRouter, Depends, HTTPException, Request
from pydantic import BaseModel, Field
from starlette.responses import Response

from llink_api.database import get_db
from llink_api.security import require_authenticated_user
from llink_api.services.security_limits import enforce_endpoint_quota

logger = logging.getLogger(__name__)
router = APIRouter()

MAX_TOOL_CODE_CHARS = 24_000
MAX_EXPORT_CHARS = 700_000
TOOL_OUTPUT_CHARS = 16_000
_SAFE_OUTPUT_NAME = re.compile(r"[^\w. -]", re.UNICODE)


class WebSearchRequest(BaseModel):
    query: str = Field(min_length=3, max_length=800)
    max_results: int = Field(default=5, ge=1, le=8)


class PythonExecutionRequest(BaseModel):
    code: str = Field(min_length=1, max_length=MAX_TOOL_CODE_CHARS)


class ExportRequest(BaseModel):
    filename: str = Field(min_length=1, max_length=160)
    content: str = Field(max_length=MAX_EXPORT_CHARS)


def safe_output_filename(value: str, extension: str) -> str:
    """Ne conserve qu'un nom de fichier simple, sans chemin ni entête HTTP injectable."""
    name = (value or "").replace("\\", "/").split("/")[-1].strip()
    name = _SAFE_OUTPUT_NAME.sub("_", name).strip(" .")
    if not name:
        name = "document"
    suffix = "." + extension
    if not name.lower().endswith(suffix):
        name = re.sub(r"\.[A-Za-z0-9]{1,8}$", "", name) + suffix
    return name[:160]


def _require_service_key(name: str) -> str:
    value = os.getenv(name, "").strip()
    if not value:
        raise HTTPException(
            status_code=503,
            detail=f"Cet outil n'est pas encore configuré sur le serveur ({name}).",
        )
    return value


async def _quota(request: Request, user: dict, db: Any, scope: str) -> None:
    await enforce_endpoint_quota(request=request, user=user, db=db, scope=scope)


def _plain_pdf_blocks(content: str) -> list[Any]:
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.platypus import Paragraph, Preformatted, Spacer

    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "LlinkBody", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=9.5, leading=14, spaceAfter=7, alignment=TA_LEFT,
    )
    code_style = ParagraphStyle(
        "LlinkCode", parent=body, fontName="Courier", fontSize=7.5,
        leading=10, leftIndent=8, rightIndent=8, backColor="#f2f3f5",
    )
    blocks: list[Any] = []
    in_code = False
    code_lines: list[str] = []
    for raw in content.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            if in_code:
                blocks.append(Preformatted("\n".join(code_lines), code_style))
                blocks.append(Spacer(1, 7))
                code_lines = []
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
        elif not line.strip():
            blocks.append(Spacer(1, 5))
        elif re.match(r"^#{1,6}\s+", line):
            level = len(line) - len(line.lstrip("#"))
            style_name = "Title" if level == 1 else "Heading2" if level == 2 else "Heading3"
            blocks.append(Paragraph(escape(re.sub(r"^#{1,6}\s+", "", line)), styles[style_name]))
        elif re.match(r"^\s*[-*+]\s+", line):
            item = re.sub(r"^\s*[-*+]\s+", "", line)
            blocks.append(Paragraph("•&nbsp;&nbsp;" + escape(item), body))
        else:
            blocks.append(Paragraph(escape(line), body))
    if code_lines:
        blocks.append(Preformatted("\n".join(code_lines), code_style))
    return blocks


def create_export_bytes(file_format: str, content: str) -> tuple[bytes, str]:
    """Génère les formats binaires depuis un texte, CSV ou JSON fourni par l'utilisateur."""
    if len(content) > MAX_EXPORT_CHARS:
        raise HTTPException(status_code=413, detail="Le contenu à exporter est trop volumineux.")

    if file_format == "pdf":
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate

        output = io.BytesIO()
        document = SimpleDocTemplate(
            output, pagesize=A4, rightMargin=20 * mm, leftMargin=20 * mm,
            topMargin=18 * mm, bottomMargin=18 * mm,
            title="Document créé avec Llink Work",
        )
        document.build(_plain_pdf_blocks(content))
        return output.getvalue(), "application/pdf"

    if file_format == "docx":
        from docx import Document
        from docx.shared import Cm, Pt

        document = Document()
        section = document.sections[0]
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        normal = document.styles["Normal"]
        normal.font.name = "Aptos"
        normal.font.size = Pt(10)
        in_code = False
        for raw in content.splitlines():
            line = raw.rstrip()
            if line.strip().startswith("```"):
                in_code = not in_code
                continue
            if in_code:
                paragraph = document.add_paragraph()
                run = paragraph.add_run(line)
                run.font.name = "Consolas"
                run.font.size = Pt(8)
            elif not line.strip():
                document.add_paragraph("")
            elif re.match(r"^#{1,6}\s+", line):
                level = min(3, len(line) - len(line.lstrip("#")))
                document.add_heading(re.sub(r"^#{1,6}\s+", "", line), level=level)
            elif re.match(r"^\s*[-*+]\s+", line):
                document.add_paragraph(re.sub(r"^\s*[-*+]\s+", "", line), style="List Bullet")
            elif re.match(r"^\s*\d+[.)]\s+", line):
                document.add_paragraph(re.sub(r"^\s*\d+[.)]\s+", "", line), style="List Number")
            else:
                document.add_paragraph(line)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    if file_format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        rows: list[list[Any]]
        stripped = content.strip()
        try:
            parsed = json.loads(stripped)
        except (json.JSONDecodeError, TypeError):
            parsed = None
        if isinstance(parsed, list) and parsed and all(isinstance(item, dict) for item in parsed):
            headers = list(dict.fromkeys(key for row in parsed for key in row))
            rows = [headers] + [[row.get(key, "") for key in headers] for row in parsed]
        elif isinstance(parsed, list) and (not parsed or all(isinstance(item, list) for item in parsed)):
            rows = parsed
        elif isinstance(parsed, dict):
            data_rows = parsed.get("rows") if isinstance(parsed.get("rows"), list) else None
            if data_rows and all(isinstance(item, dict) for item in data_rows):
                headers = list(dict.fromkeys(key for row in data_rows for key in row))
                rows = [headers] + [[row.get(key, "") for key in headers] for row in data_rows]
            else:
                rows = [["Champ", "Valeur"]] + [[key, value] for key, value in parsed.items()]
        else:
            sample = stripped[:2048]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
                rows = list(csv.reader(io.StringIO(content), dialect))
            except csv.Error:
                rows = [[line] for line in content.splitlines()]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Données"
        for row in rows:
            safe_row = []
            for value in row:
                if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
                    value = "'" + value
                safe_row.append(value if isinstance(value, (str, int, float, bool, type(None))) else json.dumps(value, ensure_ascii=False))
            sheet.append(safe_row)
        if sheet.max_row:
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="334155")
                cell.alignment = Alignment(vertical="center")
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            values = [len(str(cell.value or "")) for cell in list(column)[:200]]
            width = min(42, max(12, (max(values) if values else 10) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    raise HTTPException(status_code=422, detail="Format non pris en charge.")


@router.get("/api/work/tools")
async def list_work_tools(user: dict = Depends(require_authenticated_user)):
    """Indique seulement les outils réellement configurés côté serveur."""
    return {
        "tools": {
            "web_search": bool(os.getenv("TAVILY_API_KEY", "").strip()),
            "python": bool(os.getenv("E2B_API_KEY", "").strip()),
            "exports": True,
            "file_attachments": True,
        }
    }


@router.post("/api/work/tools/search")
async def search_web(
    payload: WebSearchRequest,
    request: Request,
    user: dict = Depends(require_authenticated_user),
    db: Any | None = Depends(get_db),
):
    await _quota(request, user, db, "work_tool_read")
    api_key = _require_service_key("TAVILY_API_KEY")
    try:
        async with httpx.AsyncClient(timeout=httpx.Timeout(25.0, connect=8.0)) as client:
            response = await client.post(
                "https://api.tavily.com/search",
                headers={"Authorization": f"Bearer {api_key}"},
                json={
                    "query": payload.query,
                    "search_depth": "basic",
                    "max_results": payload.max_results,
                    "include_answer": False,
                    "include_raw_content": False,
                },
            )
        response.raise_for_status()
        data = response.json()
        results = []
        for item in data.get("results", [])[:payload.max_results]:
            url = str(item.get("url") or "")
            if not url.startswith(("https://", "http://")):
                continue
            results.append({
                "title": str(item.get("title") or url)[:300],
                "url": url[:2000],
                "snippet": str(item.get("content") or "")[:1800],
                "published_date": item.get("published_date"),
                "score": item.get("score"),
            })
        return {"query": payload.query, "results": results}
    except HTTPException:
        raise
    except httpx.HTTPStatusError as exc:
        logger.warning("Tavily a refusé la recherche Work : HTTP %s", exc.response.status_code)
        raise HTTPException(status_code=502, detail="Le service de recherche Web a refusé la requête.") from exc
    except httpx.HTTPError as exc:
        logger.warning("Recherche Web Work indisponible : %s", exc)
        raise HTTPException(status_code=502, detail="Le service de recherche Web est indisponible.") from exc


@router.post("/api/work/tools/python")
async def run_python(
    payload: PythonExecutionRequest,
    request: Request,
    user: dict = Depends(require_authenticated_user),
    db: Any | None = Depends(get_db),
):
    """Exécute du Python dans un bac à sable E2B éphémère, jamais sur le processus Render."""
    await _quota(request, user, db, "work_tool_run")
    _require_service_key("E2B_API_KEY")
    try:
        from e2b_code_interpreter import AsyncSandbox
    except ImportError as exc:
        logger.exception("Le SDK E2B n'est pas installé.")
        raise HTTPException(status_code=503, detail="Le moteur Python isolé n'est pas installé sur le serveur.") from exc

    sandbox = None
    try:
        sandbox = await asyncio.wait_for(
            AsyncSandbox.create(timeout=60, allow_internet_access=False),
            timeout=20,
        )
        execution = await asyncio.wait_for(
            sandbox.run_code(payload.code, timeout=20),
            timeout=28,
        )
        logs = getattr(execution, "logs", None)
        stdout = "\n".join(getattr(logs, "stdout", []) or []) if logs else ""
        stderr = "\n".join(getattr(logs, "stderr", []) or []) if logs else ""
        result_text = str(getattr(execution, "text", "") or "")
        error = getattr(execution, "error", None)
        error_text = ""
        if error:
            error_text = str(getattr(error, "value", None) or getattr(error, "name", None) or error)
        return {
            "stdout": stdout[:TOOL_OUTPUT_CHARS],
            "stderr": stderr[:TOOL_OUTPUT_CHARS],
            "result": result_text[:TOOL_OUTPUT_CHARS],
            "error": error_text[:TOOL_OUTPUT_CHARS],
            "executed": True,
        }
    except asyncio.TimeoutError as exc:
        raise HTTPException(status_code=504, detail="Le temps maximal d'exécution Python a été dépassé.") from exc
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Exécution Python Work échouée.")
        raise HTTPException(status_code=502, detail="Le bac à sable Python n'a pas pu exécuter ce code.") from exc
    finally:
        if sandbox is not None:
            try:
                await asyncio.wait_for(sandbox.kill(), timeout=5)
            except Exception:
                logger.warning("Impossible de fermer immédiatement le bac à sable E2B.", exc_info=True)


@router.post("/api/work/exports/{file_format}")
async def export_work_file(
    file_format: str,
    payload: ExportRequest,
    request: Request,
    user: dict = Depends(require_authenticated_user),
    db: Any | None = Depends(get_db),
):
    await _quota(request, user, db, "work_write")
    extension = file_format.strip().lower()
    if extension not in {"pdf", "docx", "xlsx"}:
        raise HTTPException(status_code=422, detail="Choisis PDF, DOCX ou XLSX.")
    try:
        content, media_type = await asyncio.to_thread(
            create_export_bytes, extension, payload.content
        )
    except ImportError as exc:
        logger.exception("Dépendance de génération de document absente.")
        raise HTTPException(status_code=503, detail="Le moteur de génération documentaire n'est pas installé.") from exc
    filename = safe_output_filename(payload.filename, extension)
    encoded = quote(filename)
    return Response(
        content=content,
        media_type=media_type,
        headers={
            "Content-Disposition": f"attachment; filename=\"{filename.encode('ascii', 'ignore').decode('ascii')}\"; filename*=UTF-8''{encoded}",
            "X-Content-Type-Options": "nosniff",
            "Cache-Control": "no-store",
        },
    )
